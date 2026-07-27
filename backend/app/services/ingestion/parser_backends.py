from __future__ import annotations

import zipfile
from pathlib import Path

from app.services.ingestion.types import IngestionRequest, ParsedDocument, ParsedPage, TextItem
from app.services.ocr.provider import get_ocr_provider


def _title_for(request: IngestionRequest) -> str:
    return Path(request.file_name or "").stem or "未命名文档"


def _single_page_document(
    request: IngestionRequest,
    text: str,
    parser_name: str,
    parser_backend: str,
    resource_type: str | None = None,
    quality: dict | None = None,
) -> ParsedDocument:
    text = (text or "").strip()
    page = ParsedPage(
        page_number=1,
        text=text,
        text_items=[
            TextItem(
                text=text,
                page_number=1,
                confidence=1.0 if text else 0.0,
                source=parser_name,
            )
        ]
        if text
        else [],
        confidence=1.0 if text else 0.0,
        warnings=[] if text else ["empty_document"],
    )
    merged_quality = {
        "empty_pages": 0 if text else 1,
        "parser": parser_name,
        "parser_backend": parser_backend,
    }
    merged_quality.update(quality or {})
    return ParsedDocument(
        job_id=request.job_id,
        title=_title_for(request),
        source_type=request.source_type,
        resource_type=resource_type or request.resource_type,
        parser_name=parser_name,
        parser_backend=parser_backend,
        fallback_used=parser_backend != "liteparse",
        pages=[page],
        quality=merged_quality,
        metadata={"file_name": request.file_name, "content_type": request.content_type or ""},
    )


class TextFallbackParser:
    parser_name = "fallback_text"
    parser_backend = "text_fallback"

    def parse(self, path: Path, request: IngestionRequest) -> ParsedDocument:
        text = path.read_text(encoding="utf-8", errors="ignore").strip()
        resource_type = "markdown" if Path(request.file_name or "").suffix.lower() in {".md", ".markdown"} else "text"
        return _single_page_document(request, text, self.parser_name, self.parser_backend, resource_type=resource_type)


class PdfFallbackParser:
    parser_name = "fallback_pymupdf"
    parser_backend = "pdf_fallback"

    def parse(self, path: Path, request: IngestionRequest) -> ParsedDocument:
        import fitz

        pages = []
        with fitz.open(path) as pdf:
            for index, page in enumerate(pdf, start=1):
                text = page.get_text("text").strip()
                rect = page.rect
                pages.append(
                    ParsedPage(
                        page_number=index,
                        width=float(rect.width),
                        height=float(rect.height),
                        text=text,
                        text_items=[
                            TextItem(
                                text=text,
                                page_number=index,
                                width=float(rect.width),
                                height=float(rect.height),
                                confidence=1.0 if text else 0.0,
                                source=self.parser_name,
                            )
                        ]
                        if text
                        else [],
                        confidence=1.0 if text else 0.0,
                        warnings=[] if text else ["empty_page"],
                    )
                )

        return ParsedDocument(
            job_id=request.job_id,
            title=_title_for(request),
            source_type=request.source_type,
            resource_type="document",
            parser_name=self.parser_name,
            parser_backend=self.parser_backend,
            fallback_used=True,
            pages=pages,
            quality={
                "empty_pages": sum(1 for page in pages if not page.text.strip()),
                "parser": self.parser_name,
                "parser_backend": self.parser_backend,
            },
            metadata={"file_name": request.file_name, "content_type": request.content_type or ""},
        )


class OfficeTextFallbackParser:
    parser_backend = "office_text_fallback"

    def parse(self, path: Path, request: IngestionRequest) -> ParsedDocument:
        ext = Path(request.file_name or "").suffix.lower().lstrip(".")
        if ext in {"doc", "docx"}:
            parser_name = "fallback_docx"
            text = self._extract_docx_text(path)
        else:
            parser_name = "fallback_pptx"
            text = self._extract_pptx_text(path)
        return _single_page_document(request, text, parser_name, self.parser_backend, resource_type="document")

    def _extract_docx_text(self, path: Path) -> str:
        try:
            from docx import Document

            doc = Document(path)
            return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text).strip()
        except Exception:
            return path.read_text(encoding="utf-8", errors="ignore").strip()

    def _extract_pptx_text(self, path: Path) -> str:
        texts: list[str] = []
        try:
            with zipfile.ZipFile(path) as archive:
                for name in sorted(archive.namelist()):
                    if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                        xml = archive.read(name).decode("utf-8", errors="ignore")
                        texts.extend(self._extract_xml_text(xml))
        except zipfile.BadZipFile:
            return path.read_text(encoding="utf-8", errors="ignore").strip()
        return "\n".join(texts).strip()

    def _extract_xml_text(self, xml: str) -> list[str]:
        import re

        return [item.strip() for item in re.findall(r"<a:t>(.*?)</a:t>", xml) if item.strip()]


class ImageOcrParser:
    parser_name = "ocr_image"
    parser_backend = "image_ocr"

    def parse(self, path: Path, request: IngestionRequest) -> ParsedDocument:
        result = get_ocr_provider().extract_text(path)
        if isinstance(result, str):
            text = result
            confidence = 1.0 if text.strip() else 0.0
        else:
            text = result.get("text", "")
            confidence = result.get("confidence")
            confidence = float(confidence) if confidence is not None else (1.0 if text.strip() else 0.0)
        return _single_page_document(
            request,
            text,
            self.parser_name,
            self.parser_backend,
            resource_type="image",
            quality={"ocr_used": True, "ocr_confidence": confidence},
        )
